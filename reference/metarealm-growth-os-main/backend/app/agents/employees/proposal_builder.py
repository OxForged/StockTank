"""Proposal Builder — turns a deal into a real Word document.

Recipe: gather the deal, the company, and proof numbers from the
knowledge base, ask the model for the proposal sections as JSON, then
build a clean .docx with python-docx and save it into uploads/proposals.
The deal panel shows a download link. You review it before it ever
leaves the building.
"""

import uuid
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from sqlalchemy.orm import Session

from app.agents.runtime import (
    format_mad,
    humanize,
    now_iso,
    parse_json_object,
    provider_for,
)
from app.core.config import settings
from app.models import Company, Proposal
from app.models.business import Opportunity
from app.services import rag
from app.services.llm import generate, load_prompt, render_prompt

ACCENT = RGBColor(0xE1, 0x1D, 0x48)

FALLBACK_PROMPT = """Write the sections of a short sponsorship proposal in simple English. No dashes, no underscores, commas and periods only. Use only facts from the context. Answer only with JSON: {"title": "...", "intro": "...", "offer": ["...", "..."], "proof": ["...", "..."], "next_steps": "..."}

Deal context:
{context}

Proof material:
{sources}"""


def _write_docx(path: Path, company: str, data: dict, investment: str) -> None:
    doc = Document()

    heading = doc.add_heading(data["title"], level=0)
    for run in heading.runs:
        run.font.color.rgb = ACCENT

    meta = doc.add_paragraph(
        f"Prepared by MetaRealm for {company}, {datetime.now().strftime('%d %B %Y')}"
    )
    meta.runs[0].font.size = Pt(10)

    doc.add_paragraph(data["intro"])

    section = doc.add_heading("What you get", level=1)
    for run in section.runs:
        run.font.color.rgb = ACCENT
    for item in data["offer"]:
        doc.add_paragraph(str(item), style="List Bullet")

    section = doc.add_heading("Why it works", level=1)
    for run in section.runs:
        run.font.color.rgb = ACCENT
    for item in data["proof"]:
        doc.add_paragraph(str(item), style="List Bullet")

    section = doc.add_heading("Investment", level=1)
    for run in section.runs:
        run.font.color.rgb = ACCENT
    invest = doc.add_paragraph()
    invest_run = invest.add_run(investment)
    invest_run.bold = True
    invest_run.font.size = Pt(13)

    section = doc.add_heading("Next steps", level=1)
    for run in section.runs:
        run.font.color.rgb = ACCENT
    doc.add_paragraph(data["next_steps"])

    footer = doc.add_paragraph("MetaRealm, gaming marketing agency, Morocco. Lunaris Esports.")
    footer.runs[0].font.size = Pt(9)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def run(db: Session, payload: dict) -> dict:
    opportunity_id = payload.get("opportunity_id")
    if not opportunity_id:
        return {"status": "error", "summary": "Tell the Proposal Builder which deal to write for."}
    deal = db.get(Opportunity, opportunity_id)
    if deal is None:
        return {"status": "error", "summary": "Deal not found."}
    company = db.get(Company, deal.company_id)

    lines = [
        f"Deal: {deal.title} for {deal.company}, value {format_mad(deal.value_mad)}, stage {deal.stage}.",
        f"Next step on our side: {deal.next_action}.",
    ]
    if company and company.industry:
        lines.append(f"Company industry: {company.industry}.")
    if company and company.reason_to_contact:
        lines.append(f"Why them: {company.reason_to_contact}")

    sources = "\n\n".join(
        f"(From {p['document_title']}"
        + (f", page {p['page']}" if p["page"] else "")
        + f"): {p['text'][:400]}"
        for p in rag.search(db, f"{deal.title} {deal.company} results tiers numbers", top_k=4)
    ) or "No knowledge passages found."

    template = load_prompt("proposal-builder", FALLBACK_PROMPT)
    raw = generate(
        render_prompt(template, context="\n".join(lines), sources=sources),
        provider=provider_for("proposal-builder"),
    )
    if raw is None:
        return {
            "status": "llm_unavailable",
            "summary": "Could not reach the model. Start Ollama and run again.",
        }

    parsed = parse_json_object(raw) or {}
    data = {
        "title": humanize(str(parsed.get("title") or f"{deal.title}, {deal.company}"))[:120],
        "intro": humanize(str(parsed.get("intro") or raw))[:1500],
        "offer": [humanize(str(x)) for x in (parsed.get("offer") or ["Details to be discussed together."])][:6],
        "proof": [humanize(str(x)) for x in (parsed.get("proof") or ["Numbers available in the Lunaris deck."])][:5],
        "next_steps": humanize(str(parsed.get("next_steps") or "A short call this week to align on scope and timing.")),
    }
    investment = f"{format_mad(deal.value_mad)}, {deal.title}"

    safe = "".join(ch if ch.isalnum() else "-" for ch in deal.company.lower()).strip("-")
    filename = f"proposal-{safe}-{datetime.now().strftime('%Y%m%d')}-{uuid.uuid4().hex[:4]}.docx"
    path = Path(settings.UPLOADS_DIR) / "proposals" / filename
    _write_docx(path, deal.company, data, investment)

    proposal = Proposal(
        id=f"prop-{uuid.uuid4().hex[:8]}",
        opportunity_id=deal.id,
        company=deal.company,
        title=data["title"],
        filename=filename,
        created_at=now_iso(),
    )
    db.add(proposal)
    db.commit()

    return {
        "status": "completed",
        "summary": f"Proposal Builder wrote a proposal for {deal.company}: {data['title']}.",
        "details": {"proposalId": proposal.id, "filename": filename},
    }
